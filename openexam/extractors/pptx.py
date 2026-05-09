from __future__ import annotations

from pathlib import Path

from openexam.extractors.base import ExtractionError
from openexam.models import ExtractedSection


def extract_pptx(path: Path) -> list[ExtractedSection]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ExtractionError("python-pptx is not installed. Install python-pptx to parse PPTX files.") from exc

    sections: list[ExtractedSection] = []
    try:
        presentation = Presentation(path)
        for slide_index, slide in enumerate(presentation.slides, start=1):
            text_parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        text_parts.append(text)
            text = "\n".join(text_parts).strip()
            if not text:
                continue
            sections.append(
                ExtractedSection(
                    source_path=path,
                    file_name=path.name,
                    location_type="slide",
                    location_label=f"slide.{slide_index}",
                    slide_number=slide_index,
                    text=text,
                )
            )
    except Exception as exc:
        raise ExtractionError(f"Failed to parse PPTX: {exc}") from exc
    return sections
